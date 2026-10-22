"""Turn file bytes into per-page text. Uses the PDF text layer when it exists and OCR only when it doesn't."""

from __future__ import annotations

import io
import logging
from dataclasses import dataclass, field

from app.core.config import get_settings
from app.ingestion.normalize import fix_reversed_lines, looks_reversed, normalize_text
from app.ingestion.ocr import ocr_page

log = logging.getLogger(__name__)

PDF = "application/pdf"
DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
IMAGE_TYPES = {"image/png", "image/jpeg", "image/jpg", "image/tiff", "image/webp", "image/bmp"}
TEXT_TYPES = {"text/plain", "text/markdown"}
SUPPORTED_TYPES = {PDF, DOCX, *IMAGE_TYPES, *TEXT_TYPES}

OCR_DPI = 200
DOCX_CHARS_PER_PAGE = 3000  # Word files have no fixed pages; this gives stable approximate page numbers


@dataclass
class Page:
    number: int
    text: str
    source: str = "text"  # text | tesseract | vision | none
    confidence: float = 100.0


@dataclass
class Extracted:
    pages: list[Page] = field(default_factory=list)

    @property
    def text(self) -> str:
        return "\n\n".join(p.text for p in self.pages if p.text)

    @property
    def ocr_pages(self) -> list[int]:
        return [p.number for p in self.pages if p.source in ("tesseract", "vision")]

    def low_confidence_pages(self, threshold: float) -> list[int]:
        return [p.number for p in self.pages if p.source != "text" and p.confidence < threshold]


def guess_mime(name: str, mime: str | None) -> str:
    if mime and mime in SUPPORTED_TYPES:
        return mime
    lower = name.lower()
    for ext, value in (
        (".pdf", PDF),
        (".docx", DOCX),
        (".png", "image/png"),
        (".jpg", "image/jpeg"),
        (".jpeg", "image/jpeg"),
        (".tif", "image/tiff"),
        (".tiff", "image/tiff"),
        (".webp", "image/webp"),
        (".txt", "text/plain"),
        (".md", "text/markdown"),
    ):
        if lower.endswith(ext):
            return value
    return mime or "application/octet-stream"


def _clean(raw: str) -> str:
    text = normalize_text(raw)
    if looks_reversed(text):
        text = normalize_text(fix_reversed_lines(text))
    return text


def extract_pdf(data: bytes, *, allow_vision: bool = True) -> Extracted:
    import pymupdf as fitz

    s = get_settings()
    out = Extracted()
    with fitz.open(stream=data, filetype="pdf") as doc:
        for index, page in enumerate(doc, start=1):
            text = _clean(page.get_text("text", sort=True))
            useful = sum(1 for c in text if c.isalnum())
            if useful >= s.ocr_min_chars_per_page:
                out.pages.append(Page(index, text))
                continue
            png = page.get_pixmap(dpi=OCR_DPI).tobytes("png")
            result = ocr_page(png, allow_vision=allow_vision)
            out.pages.append(Page(index, _clean(result.text), result.engine, result.confidence))
    return out


def extract_image(data: bytes, *, allow_vision: bool = True) -> Extracted:
    from PIL import Image

    with Image.open(io.BytesIO(data)) as img:
        frames = []
        for i in range(getattr(img, "n_frames", 1)):  # multi-page TIFF
            img.seek(i)
            buf = io.BytesIO()
            img.convert("RGB").save(buf, format="PNG")
            frames.append(buf.getvalue())
    out = Extracted()
    for number, png in enumerate(frames, start=1):
        result = ocr_page(png, allow_vision=allow_vision)
        out.pages.append(Page(number, _clean(result.text), result.engine, result.confidence))
    return out


def extract_docx(data: bytes) -> Extracted:
    from docx import Document as DocxDocument

    doc = DocxDocument(io.BytesIO(data))
    blocks: list[str] = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                blocks.append(" | ".join(dict.fromkeys(cells)))
    return paginate("\n".join(blocks))


def paginate(raw: str) -> Extracted:
    text = _clean(raw)
    out = Extracted()
    page, size, number = [], 0, 1
    for line in text.split("\n"):
        page.append(line)
        size += len(line)
        if size >= DOCX_CHARS_PER_PAGE:
            out.pages.append(Page(number, "\n".join(page)))
            page, size, number = [], 0, number + 1
    if page:
        out.pages.append(Page(number, "\n".join(page)))
    return out


def extract(data: bytes, *, name: str, mime: str | None, allow_vision: bool = True) -> Extracted:
    kind = guess_mime(name, mime)
    if kind == PDF:
        return extract_pdf(data, allow_vision=allow_vision)
    if kind == DOCX:
        return extract_docx(data)
    if kind in IMAGE_TYPES:
        return extract_image(data, allow_vision=allow_vision)
    if kind in TEXT_TYPES:
        return paginate(data.decode("utf-8", errors="replace"))
    raise ValueError(f"סוג קובץ לא נתמך: {kind}")
