"""Render a ContractDraft to a right-to-left Hebrew Word document."""

from __future__ import annotations

import io

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

from app.agents.schemas import ContractDraft

FONT = "David"
DISCLAIMER = 'טיוטה זו הוכנה בסיוע מערכת ממוחשבת ומיועדת לעיון בלבד. יש להעבירה לבדיקת עו"ד לפני חתימה.'


def _rtl_paragraph(paragraph) -> None:
    ppr = paragraph._p.get_or_add_pPr()
    bidi = OxmlElement("w:bidi")
    bidi.set(qn("w:val"), "1")
    ppr.append(bidi)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT


def _rtl_run(run, size: float = 12, bold: bool = False, color: RGBColor | None = None) -> None:
    run.font.name = FONT
    run.font.size = Pt(size)
    run.bold = bold
    if color is not None:
        run.font.color.rgb = color
    rpr = run._r.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.append(fonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs"):
        fonts.set(qn(attr), FONT)
    rtl = OxmlElement("w:rtl")
    rtl.set(qn("w:val"), "1")
    rpr.append(rtl)
    cs_size = OxmlElement("w:szCs")
    cs_size.set(qn("w:val"), str(int(size * 2)))
    rpr.append(cs_size)
    if bold:
        rpr.append(OxmlElement("w:bCs"))


def _para(doc, text: str, *, size: float = 12, bold: bool = False, center: bool = False, space_after: float = 6,
          indent_cm: float = 0.0, color: RGBColor | None = None):
    p = doc.add_paragraph()
    _rtl_paragraph(p)
    if center:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.25
    if indent_cm:
        from docx.shared import Cm

        p.paragraph_format.right_indent = Cm(indent_cm)
    _rtl_run(p.add_run(text), size=size, bold=bold, color=color)
    return p


def build_docx(draft: ContractDraft) -> bytes:
    doc = Document()
    section = doc.sections[0]
    sect_pr = section._sectPr
    bidi = OxmlElement("w:bidi")
    sect_pr.append(bidi)

    _para(doc, draft.title, size=18, bold=True, center=True, space_after=14)
    for line in draft.opening.split("\n"):
        if line.strip():
            _para(doc, line.strip(), center=True)

    if draft.recitals:
        _para(doc, "", space_after=2)
        for r in draft.recitals:
            text = r if r.startswith("הואיל") else f"הואיל ו{r}"
            _para(doc, text, indent_cm=0.5)
        _para(doc, "לפיכך הוצהר, הותנה והוסכם בין הצדדים כדלקמן:", bold=True, space_after=10)

    for sec in draft.sections:
        _para(doc, f"{sec.number}. {sec.title}", size=13, bold=True, space_after=4)
        for clause in sec.clauses:
            depth = clause.number.count(".")
            _para(doc, f"{clause.number}  {clause.text}", indent_cm=0.6 * max(depth, 1))

    if draft.signatures:
        _para(doc, "", space_after=18)
        _para(doc, "ולראיה באו הצדדים על החתום:", bold=True, space_after=24)
        table = doc.add_table(rows=2, cols=len(draft.signatures))
        table.alignment = 2
        for i, name in enumerate(draft.signatures):
            for row, text in ((0, "_______________________"), (1, name)):
                cell = table.cell(row, i)
                p = cell.paragraphs[0]
                _rtl_paragraph(p)
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                _rtl_run(p.add_run(text))

    _para(doc, "", space_after=12)
    _para(doc, DISCLAIMER, size=9, color=RGBColor(0x80, 0x80, 0x80))

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
