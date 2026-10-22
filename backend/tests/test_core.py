"""Unit tests that need no database, network or API keys."""

from __future__ import annotations

import io

from app.agents.schemas import Clause, ContractDraft, Revision, Section
from app.core.config import Settings
from app.drive.client import DriveFile
from app.ingestion.chunker import chunk_pages
from app.ingestion.extract import Page, extract, paginate
from app.ingestion.normalize import looks_reversed, normalize_text, search_terms
from app.services.chat import cache_key, heuristic_route, web_cards
from app.services.docx_builder import build_docx
from app.services.drafting import apply_revision, render_text
from app.services.sync import plan_sync

CONTRACT_P1 = """הסכם שירותי ייעוץ
שנערך ונחתם בתל אביב ביום 01.02.2024
בין: חברת אלפא בע"מ (להלן: "המזמין")
לבין: יועצים בע"מ (להלן: "היועץ")
הואיל והמזמין מעוניין לקבל שירותי ייעוץ;
לפיכך הוסכם בין הצדדים כדלקמן:
1. מבוא ופרשנות
1.1 המבוא להסכם זה מהווה חלק בלתי נפרד ממנו.
2. השירותים
2.1 היועץ יעניק למזמין שירותי ייעוץ אסטרטגי כמפורט בנספח א'.
(א) השירותים יינתנו בהיקף של עד 40 שעות בחודש.
סודיות
3. היועץ מתחייב לשמור בסודיות מוחלטת כל מידע שיגיע אליו בקשר עם ההסכם, לתקופה של 7 שנים מתום ההסכם."""

CONTRACT_P2 = """4. התמורה
4.1 בתמורה לשירותים ישלם המזמין ליועץ סך של 15,000 ש"ח בתוספת מע"מ לחודש."""


def test_normalize_removes_niqqud_bidi_and_gershayim():
    raw = "‏שָׁלוֹם  עו״ד‎"
    assert normalize_text(raw) == 'שלום עו"ד'


def test_detects_reversed_hebrew():
    normal = "המזמין ישלם ליועץ את התמורה בתוך שלושים ימים מיום קבלת החשבונית כפי שנקבע בהסכם זה ובנספחים לו"
    reversed_text = " ".join(w[::-1] for w in normal.split())
    assert not looks_reversed(normal)
    assert looks_reversed(reversed_text)


def test_search_terms_strip_prefixes_and_stopwords():
    terms = search_terms("מה תקופת הסודיות של היועץ?")
    assert "הסודיות" in terms and "סודיות" in terms
    assert "של" not in terms and "מה" not in terms


def test_chunker_follows_clauses_and_pages():
    chunks = chunk_pages([Page(1, CONTRACT_P1), Page(2, CONTRACT_P2)])
    numbers = [c.clause_number for c in chunks]
    assert numbers[0] is None  # preamble with the parties
    assert {"2", "3", "4"} <= set(numbers)
    by_number = {c.clause_number: c for c in chunks}
    assert by_number["3"].clause_title == "סודיות"
    assert by_number["4"].page_from == 2
    assert "(א)" in by_number["2"].content  # sub-clauses stay inside their parent
    assert all(c.token_count > 0 for c in chunks)


def test_chunker_splits_long_clauses():
    long_clause = "5. אחריות\n" + " ".join(["היועץ יהיה אחראי לכל נזק שייגרם עקב רשלנותו."] * 200)
    chunks = chunk_pages([Page(1, long_clause)], max_tokens=300)
    assert len(chunks) > 3
    assert all(c.token_count <= 330 for c in chunks)
    assert chunks[1].clause_title.startswith("אחריות")


def test_text_extraction_paginates():
    out = extract("שורה ראשונה\nשורה שנייה".encode(), name="a.txt", mime=None)
    assert out.pages[0].text.startswith("שורה ראשונה")
    assert len(paginate("א" * 7000).pages) >= 1


def test_pdf_text_layer_is_used_without_ocr():
    import pymupdf as fitz

    doc = fitz.open()
    page = doc.new_page()
    page.insert_text((72, 72), "1. Services. The consultant shall provide strategy services monthly.")
    data = doc.tobytes()
    out = extract(data, name="x.pdf", mime="application/pdf", allow_vision=False)
    assert out.pages[0].source == "text"
    assert "consultant" in out.pages[0].text


def _file(fid: str, md5: str | None, mime: str = "application/pdf", name: str = "a.pdf") -> DriveFile:
    return DriveFile(fid, name, mime, name, md5, "2024-01-01T00:00:00Z", 10, None)


def test_sync_plan_detects_new_changed_removed_and_unsupported():
    remote = [_file("a", "1"), _file("b", "2-new"), _file("c", "3"), _file("z", None, "application/zip", "x.zip")]
    known = {"a": ("1", "indexed", True), "b": ("2", "indexed", True), "d": ("4", "indexed", True),
             "e": ("5", "deleted", False)}
    plan = plan_sync(remote, known)
    assert [f.id for f in plan.unchanged] == ["a"]
    assert [f.id for f in plan.changed] == ["b"]
    assert [f.id for f in plan.new] == ["c"]
    assert plan.removed == ["d"]
    assert [f.id for f in plan.unsupported] == ["z"]


def test_sync_plan_retries_failed_documents():
    plan = plan_sync([_file("a", "1")], {"a": ("1", "failed", True)})
    assert [f.id for f in plan.changed] == ["a"]


def test_google_docs_fingerprint_uses_modified_time():
    f = _file("g", None, "application/vnd.google-apps.document", "doc")
    assert f.fingerprint.startswith("mtime:")


def test_routing_heuristics():
    assert heuristic_route("תנסח לי הסכם סודיות")["intent"] == "draft"
    assert heuristic_route("מה ההבדל בין שני ההסכמים?")["intent"] == "compare"
    assert heuristic_route("מה אומר החוק על תקופת הודעה מוקדמת?")["needs_web"] is True
    assert heuristic_route("מה תקופת הסודיות?")["needs_web"] is False


def test_cache_key_changes_with_kb_version():
    assert cache_key("שאלה", 1, None) == cache_key("  שאלה ", 1, None)
    assert cache_key("שאלה", 1, None) != cache_key("שאלה", 2, None)


def test_web_cards_extract_urls():
    cards = web_cards("- חוק החוזים קובע תום לב https://www.gov.il/he/x.\n- פסק דין https://nevo.co.il/y")
    assert [c["ref"] for c in cards] == ["w1", "w2"]
    assert cards[0]["title"] == "gov.il"


def _draft() -> ContractDraft:
    return ContractDraft(
        title="הסכם ייעוץ",
        opening="שנערך ונחתם ביום [להשלמה]",
        recitals=["המזמין מעוניין בשירותי ייעוץ"],
        sections=[
            Section(number="1", title="מבוא", clauses=[Clause(number="1.1", text="המבוא חלק מההסכם.")]),
            Section(number="2", title="תמורה", clauses=[Clause(number="2.1", text="[להשלמה]")]),
        ],
        signatures=["המזמין", "היועץ"],
    )


def test_revision_patches_only_changed_sections():
    draft = _draft()
    fixed = Section(number="2", title="תמורה", clauses=[Clause(number="2.1", text="10,000 ש\"ח לחודש")])
    added = Section(number="3", title="סודיות", clauses=[Clause(number="3.1", text="סודיות מלאה.")])
    out = apply_revision(draft, Revision(sections=[added, fixed]))
    assert [s.number for s in out.sections] == ["1", "2", "3"]
    assert out.sections[1].clauses[0].text.startswith("10,000")
    assert "2.1 10,000" in render_text(out)


def test_docx_is_rtl_and_contains_text():
    from docx import Document

    data = build_docx(_draft())
    doc = Document(io.BytesIO(data))
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "הסכם ייעוץ" in text and "1.1" in text
    assert b"w:bidi" in doc.element.xml.encode() or "bidi" in doc.paragraphs[0]._p.xml


def test_render_database_url_is_converted():
    s = Settings(database_url="postgres://u:p@host:5432/db", _env_file=None)
    assert s.sqlalchemy_url == "postgresql+psycopg://u:p@host:5432/db"


DRAFT_TEXT = """כותרת: הסכם ייעוץ הנדסי
פתיח:
שנערך ונחתם בתל אביב ביום [להשלמה]
בין: אלפא נכסים בע"מ
הואיל:
- [מקור 1] והמזמין מעוניין בשירותי ייעוץ;
- והיועץ בעל ניסיון
- [חדש] לפיכך הוסכם כדלקמן:
פרק 1: מבוא ופרשנות
1.1 [מקור 2] המבוא מהווה חלק בלתי נפרד מההסכם.
1.2 [חדש] כותרות לנוחות בלבד
והן לא ישמשו לפרשנות.
פרק 2: תמורה
2.1 [אינטרנט] 22,000 ש"ח לחודש.
חתימות: המזמין | היועץ
חסר: ח.פ. היועץ; תאריך"""


def test_parse_streamed_draft_format():
    from app.services.draft_format import parse_draft, section_text

    d = parse_draft(DRAFT_TEXT)
    assert d.title == "הסכם ייעוץ הנדסי"
    assert d.opening.startswith("שנערך") and "אלפא" in d.opening
    assert d.recitals == ["המזמין מעוניין בשירותי ייעוץ;", "היועץ בעל ניסיון"]
    assert [s.number for s in d.sections] == ["1", "2"]
    c11, c12 = d.sections[0].clauses
    assert (c11.origin, c11.source) == ("internal", 2)
    assert c12.origin == "new" and c12.text.endswith("לפרשנות.")
    assert d.sections[1].clauses[0].origin == "web"
    assert d.signatures == ["המזמין", "היועץ"]
    assert d.missing_info == ["ח.פ. היועץ", "תאריך"]
    assert parse_draft(section_text(d.sections[0])).sections[0].clauses[0].source == 2


def test_parse_partial_draft_does_not_fail():
    from app.services.draft_format import parse_draft

    d = parse_draft(DRAFT_TEXT[:120])
    assert d.title == "הסכם ייעוץ הנדסי"


def test_pasted_secrets_are_cleaned():
    s = Settings(
        openai_api_key='  "sk-abc123"\n',
        google_drive_folder_id="\u200f1AbC\u200e ",
        database_url="postgresql://u:p@h/db\n",
        _env_file=None,
    )
    assert s.openai_api_key == "sk-abc123"
    assert s.google_drive_folder_id == "1AbC"
    assert s.sqlalchemy_url == "postgresql+psycopg://u:p@h/db"
    pasted = Settings(openai_api_key="sk-proj-XYZ
LOG_LEVEL=INFO
", _env_file=None)
    assert pasted.openai_api_key == "sk-proj-XYZ"
    whole_line = Settings(openai_api_key="OPENAI_API_KEY=sk-proj-XYZ", _env_file=None)
    assert whole_line.openai_api_key == "sk-proj-XYZ"
